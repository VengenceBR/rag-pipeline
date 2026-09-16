from __future__ import annotations

import hashlib
from pathlib import Path

from rag.models import RawDocument

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".html", ".htm"}


def _doc_id(company_id: str, source_path: Path) -> str:
    return hashlib.sha1(f"{company_id}:{source_path}".encode("utf-8")).hexdigest()[:16]


def load_pdf(path: Path, company_id: str) -> RawDocument:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []
    page_map: list[tuple[int, int]] = []
    offset = 0
    for page_num, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        parts.append(text)
        page_map.append((offset, page_num))
        offset += len(text) + 1
    full_text = "\n".join(parts)
    return RawDocument(
        doc_id=_doc_id(company_id, path),
        company_id=company_id,
        source_path=str(path),
        title=path.stem,
        text=full_text,
        page_map=page_map,
    )


def load_docx(path: Path, company_id: str) -> RawDocument:
    import docx

    document = docx.Document(str(path))
    full_text = "\n".join(p.text for p in document.paragraphs)
    return RawDocument(
        doc_id=_doc_id(company_id, path),
        company_id=company_id,
        source_path=str(path),
        title=path.stem,
        text=full_text,
    )


def load_html(path: Path, company_id: str) -> RawDocument:
    from bs4 import BeautifulSoup

    raw = path.read_text(encoding="utf-8", errors="ignore")
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    title = soup.title.string.strip() if soup.title and soup.title.string else path.stem
    text = soup.get_text(separator="\n")
    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    return RawDocument(
        doc_id=_doc_id(company_id, path),
        company_id=company_id,
        source_path=str(path),
        title=title,
        text=text,
    )


def load_text(path: Path, company_id: str) -> RawDocument:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return RawDocument(
        doc_id=_doc_id(company_id, path),
        company_id=company_id,
        source_path=str(path),
        title=path.stem,
        text=text,
    )


_LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".html": load_html,
    ".htm": load_html,
    ".txt": load_text,
    ".md": load_text,
}


def load_document(path: Path, company_id: str) -> RawDocument:
    ext = path.suffix.lower()
    loader = _LOADERS.get(ext)
    if loader is None:
        raise ValueError(f"Unsupported file type: {ext} ({path})")
    return loader(path, company_id)


def load_directory(directory: Path, company_id: str) -> list[RawDocument]:
    docs = []
    for path in sorted(directory.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            docs.append(load_document(path, company_id))
    return docs
